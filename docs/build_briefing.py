# -*- coding: utf-8 -*-
"""One-page executive briefing (RTL Persian) for CEO review."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Pt, RGBColor, Emu, Twips

BLUE = RGBColor(0x15, 0x65, 0xC0)
DARK = RGBColor(0x1A, 0x23, 0x32)
MUTED = RGBColor(0x5A, 0x67, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_BG = RGBColor(0xF4, 0xF7, 0xFB)
ACCENT_BG = RGBColor(0xE8, 0xF1, 0xFB)
GREEN = RGBColor(0x1B, 0x7A, 0x4E)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def set_run(run, size=11, bold=False, color=DARK, font="Tahoma"):
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _el("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), font)
    if rPr.find(qn("w:rtl")) is None:
        rPr.append(_el("w:rtl"))
    if rPr.find(qn("w:cs")) is None:
        rPr.append(_el("w:cs"))


def set_para_rtl(p, align="right", space_before=0, space_after=4, line=1.08):
    p.alignment = {
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        pPr.append(_el("w:bidi", **{"w:val": "1"}))
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = _el("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), {"right": "right", "center": "center", "left": "left", "justify": "both"}[align])


def add_text(p, text, size=11, bold=False, color=DARK):
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return run


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = _el("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = _el("w:tcMar")
        tcPr.append(tcMar)
    for name, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = _el(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_cell_valign(cell, val="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    vAlign = tcPr.find(qn("w:vAlign"))
    if vAlign is None:
        vAlign = _el("w:vAlign")
        tcPr.append(vAlign)
    vAlign.set(qn("w:val"), val)


def set_table_borders(table, color="D0D7DE", sz="4", inside=True):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else _el("w:tblPr")
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = _el("w:tblBorders")
        tblPr.append(borders)
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in edges:
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = _el(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else _el("w:tblPr")
    borders = _el("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = _el(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    tblPr.append(borders)


def set_table_rtl(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    bidi = tblPr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = _el("w:bidiVisual")
        tblPr.append(bidi)
    bidi.set(qn("w:val"), "1")


def cell_para(cell, text, size=10, bold=False, color=DARK, align="right", space_after=0):
    cell.text = ""
    p = cell.paragraphs[0]
    set_para_rtl(p, align=align, space_before=0, space_after=space_after, line=1.12)
    add_text(p, text, size=size, bold=bold, color=color)
    return p


def heading(doc, text):
    p = doc.add_paragraph()
    set_para_rtl(p, align="right", space_before=8, space_after=3, line=1.0)
    add_text(p, text, size=12, bold=True, color=BLUE)
    # underline via bottom border on paragraph
    pPr = p._p.get_or_add_pPr()
    pBdr = _el("w:pBdr")
    bottom = _el("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1565C0")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def body(doc, text, after=4, justify=True):
    p = doc.add_paragraph()
    set_para_rtl(p, align="justify" if justify else "right", space_before=0, space_after=after, line=1.12)
    add_text(p, text, size=10.5, bold=False, color=DARK)
    return p


def bullet(doc, title, rest):
    p = doc.add_paragraph()
    set_para_rtl(p, align="right", space_before=0, space_after=2, line=1.12)
    add_text(p, "■  ", size=9, bold=True, color=BLUE)
    add_text(p, title, size=10.5, bold=True, color=DARK)
    if rest:
        add_text(p, "  —  " + rest, size=10.5, bold=False, color=DARK)
    return p


def set_document_rtl(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Tahoma"
    normal.font.size = Pt(11)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = _el("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), "Tahoma")
    # section bidi
    for section in doc.sections:
        sectPr = section._sectPr
        if sectPr.find(qn("w:bidi")) is None:
            sectPr.append(_el("w:bidi", **{"w:val": "1"}))


def build():
    doc = Document()
    set_document_rtl(doc)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    # ── Header bar ──────────────────────────────────────────────
    header = doc.add_table(rows=1, cols=1)
    header.autofit = True
    set_table_rtl(header)
    no_borders(header)
    cell = header.cell(0, 0)
    shade_cell(cell, "1565C0")
    set_cell_margins(cell, top=90, bottom=90, left=140, right=140)
    set_cell_valign(cell, "center")
    cell.text = ""
    p = cell.paragraphs[0]
    set_para_rtl(p, align="center", space_before=0, space_after=2, line=1.0)
    add_text(p, "خلاصه محصول برای بررسی مدیریت", size=11, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    set_para_rtl(p2, align="center", space_before=0, space_after=0, line=1.0)
    add_text(p2, "درخواست جلسه معرفی  ·  ۲۰ تا ۳۰ دقیقه  ·  ۲ شهریور ۱۴۰۵", size=9, bold=False, color=RGBColor(0xC5, 0xDA, 0xF5))

    # ── Title ───────────────────────────────────────────────────
    p = doc.add_paragraph()
    set_para_rtl(p, align="right", space_before=10, space_after=1, line=1.0)
    add_text(p, "مدیریت زمان شخصی", size=20, bold=True, color=BLUE)

    p = doc.add_paragraph()
    set_para_rtl(p, align="right", space_before=0, space_after=6, line=1.05)
    add_text(p, "اپلیکیشن اندروید حضور، تسک و گزارش کار  |  محصول آماده دمو  |  کاملاً آفلاین", size=10.5, bold=False, color=MUTED)

    # meta strip
    meta = doc.add_table(rows=1, cols=3)
    set_table_rtl(meta)
    no_borders(meta)
    labels = [
        ("وضعیت", "نسخه قابل نصب و ارائه"),
        ("پلتفرم", "اندروید بومی (Kotlin)"),
        ("تهیه‌کننده", "امید"),
    ]
    for i, (k, v) in enumerate(labels):
        c = meta.cell(0, i)
        shade_cell(c, "F4F7FB")
        set_cell_margins(c, top=50, bottom=50, left=70, right=70)
        set_cell_valign(c, "center")
        c.text = ""
        pk = c.paragraphs[0]
        set_para_rtl(pk, align="center", space_before=0, space_after=0, line=1.0)
        add_text(pk, k, size=8, bold=False, color=MUTED)
        pv = c.add_paragraph()
        set_para_rtl(pv, align="center", space_before=0, space_after=0, line=1.0)
        add_text(pv, v, size=10, bold=True, color=DARK)

    # ── Purpose ─────────────────────────────────────────────────
    heading(doc, "هدف از این جلسه")
    body(
        doc,
        "معرفی یک محصول نرم‌افزاری که برای مدیریت زمان کاری، ثبت حضور، پیگیری تسک‌ها و گزارش‌گیری توسعه داده شده است؛ "
        "و دریافت نظر جنابعالی درباره امکان استفاده آزمایشی، سفارشی‌سازی، یا توسعه آن در مجموعه.",
        after=2,
    )

    # ── Problem / Solution ──────────────────────────────────────
    heading(doc, "مسئله و راه‌حل")
    body(
        doc,
        "ثبت حضور، پیگیری کار روی پروژه‌ها و تهیه گزارش زمان معمولاً پراکنده و دستی است. تصویر دقیقی از ساعت کار واقعی، "
        "شناوری، مرخصی، اضافه‌کاری و زمان صرف‌شده روی هر تسک به‌راحتی در دسترس نیست.",
        after=3,
    )
    body(
        doc,
        "این اپلیکیشن این کارها را در یک ابزار بومی، ساده و بدون وابستگی به اینترنت جمع می‌کند. منطق ساعت کاری ایران "
        "(تقویم شمسی، شناوری، پنجشنبه و تعطیلات) در آن پیاده شده و داده‌ها فقط روی خود دستگاه ذخیره می‌شود.",
        after=2,
    )

    # ── Features table ──────────────────────────────────────────
    heading(doc, "قابلیت‌های اصلی")
    feat = doc.add_table(rows=5, cols=2)
    set_table_rtl(feat)
    set_table_borders(feat, color="D6DEE8", sz="4")
    rows = [
        ("حضور و خروج", "ثبت ورود و خروج، ویجت سریع روی صفحه اصلی، محاسبه مرخصی و اضافه‌کاری، پیشنهاد ساعت پایان کار، یادآوری"),
        ("تسک و پروژه", "تعریف تسک با پروژه و شماره جیرا، تایمر زنده، برآورد زمان و باقی‌مانده، ثبت لاگ کار روی هر تسک"),
        ("گزارش و تقویم", "گزارش روزانه، هفتگی و ماهانه با نمودار؛ تقویم شمسی، تعطیلات و برنامه هفتگی قابل تنظیم"),
        ("موقعیت محل کار", "امکان هشدار یا ورود/خروج خودکار هنگام رسیدن به محدوده تعریف‌شده محل کار"),
        ("امنیت و پشتیبان", "قفل اثرانگشت/چهره، پشتیبان JSON و پشتیبان خودکار؛ بدون سرور خارجی و بدون ارسال داده به بیرون"),
    ]
    for i, (title, desc) in enumerate(rows):
        c0, c1 = feat.cell(i, 0), feat.cell(i, 1)
        shade_cell(c0, "1565C0" if i == 0 else ("E8F1FB" if i % 2 == 0 else "FFFFFF"))
        shade_cell(c1, "E8F1FB" if i % 2 == 0 else "FFFFFF")
        # actually first col is always accent
        shade_cell(c0, "1565C0")
        set_cell_margins(c0, top=45, bottom=45, left=70, right=70)
        set_cell_margins(c1, top=45, bottom=45, left=80, right=80)
        set_cell_valign(c0, "center")
        set_cell_valign(c1, "center")
        cell_para(c0, title, size=10, bold=True, color=WHITE, align="center")
        cell_para(c1, desc, size=10, bold=False, color=DARK, align="right")

    # set col widths: title narrow, desc wide
    for row in feat.rows:
        row.cells[0].width = Cm(3.6)
        row.cells[1].width = Cm(14.6)

    p = doc.add_paragraph()
    set_para_rtl(p, align="right", space_before=4, space_after=2, line=1.1)
    add_text(p, "نکته تکمیلی: ", size=10, bold=True, color=BLUE)
    add_text(
        p,
        "آیکون برنامه ساعت کار امروز را نشان می‌دهد؛ حالت تاریک و تم رنگی هم قابل تنظیم است.",
        size=10,
        bold=False,
        color=DARK,
    )

    # ── Tech ────────────────────────────────────────────────────
    heading(doc, "مشخصات فنی برای بررسی کارشناسی")
    tech = doc.add_table(rows=1, cols=4)
    set_table_rtl(tech)
    no_borders(tech)
    tech_items = [
        ("فناوری", "Kotlin + Room\nMaterial 3"),
        ("حداقل سیستم", "اندروید ۸\nو بالاتر"),
        ("داده", "کاملاً محلی\nبدون سرور"),
        ("خروجی", "APK ریلیز\nآماده نصب"),
    ]
    for i, (k, v) in enumerate(tech_items):
        c = tech.cell(0, i)
        shade_cell(c, "F4F7FB")
        set_cell_margins(c, top=50, bottom=50, left=50, right=50)
        set_cell_valign(c, "center")
        c.text = ""
        pk = c.paragraphs[0]
        set_para_rtl(pk, align="center", space_before=0, space_after=1, line=1.0)
        add_text(pk, k, size=8, bold=False, color=MUTED)
        for j, line in enumerate(v.split("\n")):
            pv = c.add_paragraph()
            set_para_rtl(pv, align="center", space_before=0, space_after=0, line=1.05)
            add_text(pv, line, size=10, bold=True, color=DARK)

    # ── Ask ─────────────────────────────────────────────────────
    heading(doc, "پیشنهاد و درخواست")
    body(
        doc,
        "در صورت صلاحدید، یک جلسه ۲۰ تا ۳۰ دقیقه‌ای برای دموی زنده هماهنگ شود. در همان جلسه می‌توان این گزینه‌ها را جمع‌بندی کرد:",
        after=3,
    )
    bullet(doc, "استفاده آزمایشی در تیم", "نصب روی چند دستگاه و ارزیابی در کار واقعی")
    bullet(doc, "سفارشی‌سازی", "تنظیم مطابق شیفت، پروژه‌ها و فرآیند داخلی مجموعه")
    bullet(doc, "توسعه سازمانی", "در صورت نیاز، مسیر نسخه تیمی یا گزارش مدیریتی")

    # closing box
    close = doc.add_table(rows=1, cols=1)
    set_table_rtl(close)
    no_borders(close)
    cc = close.cell(0, 0)
    shade_cell(cc, "E8F1FB")
    set_cell_margins(cc, top=70, bottom=70, left=120, right=120)
    cc.text = ""
    p = cc.paragraphs[0]
    set_para_rtl(p, align="right", space_before=0, space_after=2, line=1.12)
    add_text(p, "آماده ارائه دمو در جلسه هستم. هدف، معرفی محصول و گرفتن نظر جنابعالی است؛ نه الزام به تصمیم فوری.", size=10.5, bold=False, color=DARK)
    p2 = cc.add_paragraph()
    set_para_rtl(p2, align="left", space_before=4, space_after=0, line=1.0)
    add_text(p2, "با احترام  ·  امید", size=10.5, bold=True, color=BLUE)

    out = r"e:\Kotlin\PersonalTimeTracker\docs\Briefing-PersonalTimeTracker.docx"
    doc.save(out)
    print("saved")


if __name__ == "__main__":
    build()
